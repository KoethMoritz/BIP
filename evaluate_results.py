#!/usr/bin/env python3
"""
Evaluation der Metadaten-Extraktion (Title, Date, ID, Scale) gegen eine Groundtruth.

Groundtruth als .xlsx (erstes Blatt) oder .csv mit den Spalten:

    File Name | Title | Date | ID | Scale | Datatype   (Datatype wird ignoriert)

CSV-Beispiel (Trenner ; oder , wird automatisch erkannt):

    File Name;Title;Date;ID;Scale
    plan_001.tif;Bahnhof Alexanderplatz Grundriss|Alexanderplatz Grundriss;12.03.1985;A-1234;1:100|1:50
    plan_002.tif;Not found;01.1990;Not found;1:200

  - Mehrere gültige Lösungen: mit "|" trennen.
  - Feld nicht vorhanden im Plan: "Not found" oder leer lassen.
  - Dateiendung im filename ist optional.

Aufruf:
    python evaluate_results.py                     # fragt beide Pfade im Terminal ab
    python evaluate_results.py --results results_api.txt --gt groundtruth.xlsx
    python evaluate_results.py --llm-judge      # Titel-Grenzfälle zusätzlich vom LLM bewerten lassen
"""

import argparse
import csv
import os
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path

FIELDS = ["Title", "Date", "ID", "Scale"]
JUDGE_FIELDS = {"Title"}  # nur hier ist ein LLM-Urteil sinnvoll (ID/Datum/Maßstab bleiben strikt)
KNOWN_EXT = (".tif", ".tiff", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".pdf")
NOT_FOUND = {"", "not found", "notfound", "n/a", "na", "none", "null", "-", "--",
             "unknown", "nicht gefunden", "nicht vorhanden", "kein", "keine"}
GT_ALIASES = {
    "filename": "filename", "file name": "filename", "file": "filename", "datei": "filename", "dateiname": "filename",
    "title": "Title", "titel": "Title",
    "date": "Date", "datum": "Date",
    "id": "ID", "nummer": "ID", "plannummer": "ID",
    "scale": "Scale", "maßstab": "Scale", "massstab": "Scale", "masstab": "Scale",
}
MONTHS = {
    "januar": 1, "jan": 1, "january": 1, "februar": 2, "feb": 2, "february": 2,
    "märz": 3, "maerz": 3, "mär": 3, "mar": 3, "march": 3, "april": 4, "apr": 4,
    "mai": 5, "may": 5, "juni": 6, "jun": 6, "june": 6, "juli": 7, "jul": 7, "july": 7,
    "august": 8, "aug": 8, "september": 9, "sep": 9, "sept": 9, "oktober": 10, "okt": 10,
    "oct": 10, "october": 10, "november": 11, "nov": 11, "dezember": 12, "dez": 12,
    "dec": 12, "december": 12,
}


# ============================================================
# Hilfsfunktionen
# ============================================================

def is_empty(value: str) -> bool:
    return value.strip().strip(".").strip().lower() in NOT_FOUND


def file_key(name: str) -> str:
    name = name.strip()
    if name.lower().endswith(KNOWN_EXT):
        name = Path(name).stem
    return name.lower()


# ============================================================
# Normalisierung
# ============================================================

def norm_title(s: str) -> str:
    s = unicodedata.normalize("NFKC", s).casefold()
    s = re.sub(r"[^\w\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_id(s: str) -> str:
    s = unicodedata.normalize("NFKC", s).casefold()
    return re.sub(r"[\s\-_/.,:;]", "", s)


def norm_date(s: str) -> str:
    """Gibt ISO-artigen String zurück (YYYY-MM-DD, YYYY-MM oder YYYY); sonst normalisierten Text."""
    t = unicodedata.normalize("NFKC", s).casefold().strip()

    def year(y: str) -> int:
        y_int = int(y)
        if len(y) == 2:
            return 2000 + y_int if y_int <= 30 else 1900 + y_int
        return y_int

    m = re.search(r"\b(\d{4})-(\d{1,2})-(\d{1,2})\b", t)
    if m:
        return f"{int(m[1]):04d}-{int(m[2]):02d}-{int(m[3]):02d}"
    m = re.search(r"\b(\d{1,2})\s*[.\-/ ]\s*(\d{1,2})\s*[.\-/ ]\s*(\d{4}|\d{2})\b", t)
    if m:
        return f"{year(m[3]):04d}-{int(m[2]):02d}-{int(m[1]):02d}"
    m = re.search(r"\b(\d{1,2})\.?\s*([a-zäöü]{3,9})\.?\s*(\d{4}|\d{2})\b", t)
    if m and m[2] in MONTHS:
        return f"{year(m[3]):04d}-{MONTHS[m[2]]:02d}-{int(m[1]):02d}"
    m = re.search(r"\b([a-zäöü]{3,9})\.?\s*(\d{4})\b", t)
    if m and m[1] in MONTHS:
        return f"{int(m[2]):04d}-{MONTHS[m[1]]:02d}"
    m = re.search(r"\b(\d{1,2})\s*[./\-]\s*(\d{4})\b", t)
    if m and 1 <= int(m[1]) <= 12:
        return f"{int(m[2]):04d}-{int(m[1]):02d}"
    m = re.search(r"\b(1[6-9]\d\d|20\d\d)\b", t)
    if m:
        return m[1]
    return re.sub(r"\s+", "", t)


def norm_scale(s: str) -> frozenset:
    """Extrahiert alle Verhältnisse a:b bzw. a/b -> {'1:100', ...}."""
    t = unicodedata.normalize("NFKC", s)
    found = {f"{int(a)}:{int(b)}" for a, b in re.findall(r"(\d+)\s*[:/]\s*(\d+)", t)}
    if found:
        return frozenset(found)
    return frozenset({re.sub(r"\s+", "", t.casefold())})


# ============================================================
# Vergleich  ->  Score 0..1
# ============================================================

def sim_title(a: str, b: str) -> float:
    na, nb = norm_title(a), norm_title(b)
    if not na or not nb:
        return 0.0
    r1 = SequenceMatcher(None, na, nb).ratio()
    r2 = SequenceMatcher(None, " ".join(sorted(na.split())), " ".join(sorted(nb.split()))).ratio()
    return max(r1, r2)


def sim_id(a: str, b: str) -> float:
    na, nb = norm_id(a), norm_id(b)
    if not na or not nb:
        return 0.0
    return SequenceMatcher(None, na, nb).ratio()


def sim_date(a: str, b: str) -> float:
    return 1.0 if norm_date(a) == norm_date(b) else 0.0


def sim_scale(a: str, b: str) -> float:
    return 1.0 if norm_scale(a) & norm_scale(b) else 0.0


SIM = {"Title": sim_title, "Date": sim_date, "ID": sim_id, "Scale": sim_scale}


# ============================================================
# Einlesen
# ============================================================

def parse_results(path: Path):
    """Liest results_api.txt. Rückgabe: {file_key: {field: [werte pro Seite]}}, Statistiken."""
    text = path.read_text(encoding="utf-8", errors="replace")
    entry_re = re.compile(r"^FILE:\s*(.+?)\n-{5,}\n(.*?)\n={5,}", re.S | re.M)
    entries, dupes, errors = {}, 0, 0

    for m in entry_re.finditer(text):
        header, body = m.group(1).strip(), m.group(2)
        hm = re.match(r"^(.*?)(?:\s*\(Page\s+(\d+)/(\d+)\))?$", header)
        fname, page = hm.group(1).strip(), int(hm.group(2) or 1)
        key = (file_key(fname), page)

        if body.strip().lower().startswith("error"):
            errors += 1
        fields = {}
        for line in body.splitlines():
            line = line.replace("*", "").strip()
            fm = re.match(r"^(Title|Date|ID|Scale)\s*:\s*(.*)$", line, re.I)
            if fm:
                canon = {f.lower(): f for f in FIELDS}[fm.group(1).lower()]
                fields[canon] = fm.group(2).strip()
        if key in entries:
            dupes += 1  # Datei wurde mehrfach verarbeitet (Append-Modus) -> letzter Lauf gewinnt
        entries[key] = fields

    by_file = defaultdict(lambda: {f: [] for f in FIELDS})
    for (fk, _page), fields in sorted(entries.items()):
        for f in FIELDS:
            v = fields.get(f, "")
            if not is_empty(v):
                by_file[fk][f].append(v)
    return by_file, {"entries": len(entries), "duplicates": dupes, "errors": errors}


def _cell_to_str(v) -> str:
    """Excel-Zelle -> String (Datums-Zellen als TT.MM.JJJJ, ganze Zahlen ohne '.0')."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    if hasattr(v, "strftime"):  # datetime.date
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_gt_rows(path: Path):
    """Liest .xlsx/.xlsm (erstes Blatt) oder .csv -> (Spaltennamen, Liste von Zeilen-Dicts)."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Für .xlsx wird openpyxl benötigt:  pip install openpyxl")
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            sys.exit("Groundtruth-Excel ist leer.")
        header = [_cell_to_str(h) for h in rows[0]]
        data = [dict(zip(header, [_cell_to_str(c) for c in r])) for r in rows[1:]]
        return header, data

    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    first = raw.splitlines()[0] if raw else ""
    delim = ";" if first.count(";") >= first.count(",") else ","
    reader = csv.DictReader(raw.splitlines(), delimiter=delim)
    return list(reader.fieldnames or []), list(reader)


def parse_groundtruth(path: Path):
    header, rows = _read_gt_rows(path)

    colmap = {}
    for col in header:
        key = re.sub(r"[\s_]+", " ", col.strip().lower())
        canon = GT_ALIASES.get(key)
        if canon and canon not in colmap:
            colmap[canon] = col  # Spalten wie "Datatype" werden ignoriert
    missing = [c for c in ["filename"] + FIELDS if c not in colmap]
    if missing:
        sys.exit(f"Groundtruth: Spalten fehlen oder sind nicht erkannt: {missing}. "
                 f"Gefunden: {header}")

    gt = {}
    for row in rows:
        fname = (row.get(colmap["filename"]) or "").strip()
        if not fname:
            continue
        entry = {}
        for f in FIELDS:
            cell = (row.get(colmap[f]) or "").strip()
            alts = [a.strip() for a in cell.split("|")]
            entry[f] = [a for a in alts if not is_empty(a)]
        gt[file_key(fname)] = entry
    return gt


# ============================================================
# Optionaler LLM-Judge (nur für Grenzfälle)
# ============================================================

def make_judge():
    from dotenv import load_dotenv
    from openai import OpenAI

    load_dotenv(Path(__file__).resolve().parent / ".env")
    client = OpenAI(api_key=os.getenv("HTW_API_KEY"), base_url=os.getenv("HTW_BASE_URL"))
    model = os.getenv("MODEL_NAME")

    def judge(field: str, gts: list, pred: str) -> bool:
        prompt = (
            f"Feld: {field}\n"
            f"Erlaubte Referenzwerte: {' | '.join(gts)}\n"
            f"Vorhergesagter Wert: {pred}\n\n"
            "Bezeichnet der vorhergesagte Wert inhaltlich dasselbe wie einer der Referenzwerte "
            "(Abkürzungen, Schreibweise, kleine OCR-Tippfehler, Wortreihenfolge sind okay; "
            "anderer Inhalt ist nicht okay)? Antworte NUR mit JA oder NEIN."
        )
        r = client.chat.completions.create(
            model=model, messages=[{"role": "user", "content": prompt}],
            temperature=0, max_tokens=300,
        )
        ans = re.sub(r"<think>.*?</think>", "", r.choices[0].message.content or "", flags=re.S)
        return bool(re.match(r"\s*(JA|YES)\b", ans.strip(), re.I))

    return judge


# ============================================================
# Auswertung
# ============================================================

def prf(tp, fp, fn):
    p = tp / (tp + fp) if tp + fp else 0.0
    r = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * p * r / (p + r) if p + r else 0.0
    return p, r, f1


def evaluate(gt, preds, title_thr, id_thr, judge=None):
    thr = {"Title": title_thr, "ID": id_thr, "Date": 1.0, "Scale": 1.0}
    counts = {f: dict(tp=0, fp=0, fn=0, tn=0) for f in FIELDS}
    details = []
    file_all_ok = {}

    for fk, g in gt.items():
        p = preds.get(fk)
        file_all_ok[fk] = True
        for f in FIELDS:
            gts = g[f]
            cands = p[f] if p else []
            best, best_pred = 0.0, ""
            for c in cands:
                for a in gts:
                    s = SIM[f](a, c)
                    if s > best:
                        best, best_pred = s, c
            judged = ""

            if not gts and not cands:
                outcome = "TN"
            elif not gts and cands:
                outcome = "FP"
            elif gts and not cands:
                outcome = "FN"
            else:
                match = best >= thr[f]
                if not match and judge and f in JUDGE_FIELDS:
                    try:
                        match = any(judge(f, gts, c) for c in cands)
                        judged = "JA" if match else "NEIN"
                    except Exception as e:  # LLM nicht erreichbar -> strikt bleiben
                        judged = f"Fehler: {e}"
                outcome = "TP" if match else "FP+FN"

            c = counts[f]
            if outcome == "TP":
                c["tp"] += 1
            elif outcome == "TN":
                c["tn"] += 1
            elif outcome == "FP":
                c["fp"] += 1
            elif outcome == "FN":
                c["fn"] += 1
            else:
                c["fp"] += 1
                c["fn"] += 1
            if outcome not in ("TP", "TN"):
                file_all_ok[fk] = False

            details.append({
                "file": fk, "field": f,
                "groundtruth": " | ".join(gts) if gts else "Not found",
                "prediction": " || ".join(cands) if cands else "Not found",
                "similarity": f"{best:.3f}" if gts and cands else "",
                "outcome": outcome, "llm_judge": judged,
                "file_in_results": "yes" if p else "NO",
            })
    return counts, details, file_all_ok


def main():
    ap = argparse.ArgumentParser(description="Evaluation OCR/LLM-Metadaten vs. Groundtruth")
    ap.add_argument("--results", default=None, help="Pfad zur Results-TXT (sonst Abfrage im Terminal)")
    ap.add_argument("--gt", default=None, help="Pfad zur Groundtruth .xlsx/.csv (sonst Abfrage im Terminal)")
    ap.add_argument("--out-dir", default="eval_output")
    ap.add_argument("--title-threshold", type=float, default=0.85,
                    help="Min. Ähnlichkeit (0-1) für Titel-Treffer (Standard 0.85)")
    ap.add_argument("--id-threshold", type=float, default=1.0,
                    help="Min. Ähnlichkeit für ID-Treffer; 1.0 = exakt nach Normalisierung")
    ap.add_argument("--llm-judge", action="store_true",
                    help="Titel-Nichttreffer zusätzlich vom lokalen LLM prüfen lassen")
    args = ap.parse_args()

    def ask_path(prompt: str) -> str:
        while True:
            val = input(prompt).strip().strip("\"'")
            if val:
                return val
            print("  Bitte einen Pfad eingeben.")

    results_arg = args.results or ask_path("Pfad zur Results-Datei (z.B. results_api.txt): ")
    gt_arg = args.gt or ask_path("Pfad zur Groundtruth-Datei (.xlsx/.csv): ")
    results_path, gt_path = Path(results_arg.strip().strip("\"'")), Path(gt_arg.strip().strip("\"'"))
    for p in (results_path, gt_path):
        if not p.exists():
            sys.exit(f"Datei nicht gefunden: {p}")

    gt = parse_groundtruth(gt_path)
    preds, stats = parse_results(results_path)
    judge = make_judge() if args.llm_judge else None

    counts, details, file_all_ok = evaluate(gt, preds, args.title_threshold, args.id_threshold, judge)

    # ---------- Kennzahlen ----------
    n_docs = len(gt)
    lines, csv_rows = [], []
    tot = dict(tp=0, fp=0, fn=0, tn=0)
    f1s, recalls, precs, accs = [], [], [], []
    for f in FIELDS:
        c = counts[f]
        correct = c["tp"] + c["tn"]
        acc = correct / n_docs if n_docs else 0.0
        p, r, f1 = prf(c["tp"], c["fp"], c["fn"])
        f1s.append(f1); recalls.append(r); precs.append(p); accs.append(acc)
        for k in tot:
            tot[k] += c[k]
        csv_rows.append([f, n_docs, c["tp"], c["fp"], c["fn"], c["tn"],
                         f"{acc:.4f}", f"{p:.4f}", f"{r:.4f}", f"{f1:.4f}"])

    # Overall (micro): über alle Felder aller Dokumente
    total_cells = n_docs * len(FIELDS)
    o_acc = (tot["tp"] + tot["tn"]) / total_cells if total_cells else 0.0
    o_p, o_r, o_f1 = prf(tot["tp"], tot["fp"], tot["fn"])
    csv_rows.append(["OVERALL (micro)", total_cells, tot["tp"], tot["fp"], tot["fn"], tot["tn"],
                     f"{o_acc:.4f}", f"{o_p:.4f}", f"{o_r:.4f}", f"{o_f1:.4f}"])
    m = len(FIELDS)
    csv_rows.append(["MACRO (Ø Felder)", "", "", "", "", "",
                     f"{sum(accs)/m:.4f}", f"{sum(precs)/m:.4f}", f"{sum(recalls)/m:.4f}", f"{sum(f1s)/m:.4f}"])
    docs_perfect = sum(1 for v in file_all_ok.values() if v)
    csv_rows.append(["ALLE 4 FELDER KORREKT (pro Dokument)", n_docs, docs_perfect, "", "", "",
                     f"{docs_perfect/n_docs:.4f}" if n_docs else "0", "", "", ""])

    header = ["Field", "N", "TP", "FP", "FN", "TN", "Accuracy", "Precision", "Recall", "F1"]

    missing_in_results = [k for k in gt if k not in preds]
    not_in_gt = [k for (k) in preds if k not in gt]

    # ---------- Ausgabe ----------
    out = []
    out.append("=" * 78)
    out.append(f"EVALUATION  ({datetime.now():%Y-%m-%d %H:%M:%S})")
    out.append(f"Results: {results_path}   |   Groundtruth: {gt_path}")
    out.append(f"Titel-Schwelle: {args.title_threshold}  |  ID-Schwelle: {args.id_threshold}  |  "
               f"LLM-Judge: {'an' if args.llm_judge else 'aus'}")
    out.append("=" * 78)
    out.append(f"Dokumente in Groundtruth: {n_docs}")
    out.append(f"Einträge in Results: {stats['entries']} "
               f"(Duplikate/Mehrfachläufe: {stats['duplicates']}, Fehlerzeilen: {stats['errors']})")
    if missing_in_results:
        out.append(f"WARNUNG: {len(missing_in_results)} GT-Dokumente fehlen in Results: "
                   f"{', '.join(missing_in_results[:10])}{' ...' if len(missing_in_results) > 10 else ''}")
    if not_in_gt:
        out.append(f"Hinweis: {len(not_in_gt)} Results-Dateien ohne GT-Eintrag (ignoriert)")
    out.append("")
    fmt = "{:<38}{:>4}{:>5}{:>5}{:>5}{:>5}{:>10}{:>10}{:>8}{:>8}"
    out.append(fmt.format("Feld", "N", "TP", "FP", "FN", "TN", "Accuracy", "Precision", "Recall", "F1"))
    out.append("-" * 98)
    for r_ in csv_rows:
        vals = [str(x) for x in r_]
        out.append(fmt.format(*vals))
    out.append("")
    out.append("Definition: TP = korrekter Wert erkannt | TN = korrekt 'Not found' | FP = Wert erkannt, "
               "der falsch ist/nicht existiert | FN = existierender Wert nicht (korrekt) erkannt.")
    out.append("Ein falscher Wert zählt als FP *und* FN.")
    summary_txt = "\n".join(out)
    print(summary_txt)

    # ---------- Speichern ----------
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    (out_dir / f"metrics_{stamp}.txt").write_text(summary_txt, encoding="utf-8")
    with open(out_dir / f"metrics_{stamp}.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(header)
        w.writerows(csv_rows)
    with open(out_dir / f"details_{stamp}.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(details[0].keys()), delimiter=";")
        w.writeheader()
        w.writerows(details)

    print(f"\nGespeichert in: {out_dir.resolve()}")
    print(f"  metrics_{stamp}.txt / .csv  (Zusammenfassung)")
    print(f"  details_{stamp}.csv         (jede Einzelentscheidung zum Nachprüfen)")


if __name__ == "__main__":
    main()